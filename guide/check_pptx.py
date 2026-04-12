from pptx import Presentation
import os

pptx_path = r"c:\Users\40270\Desktop\workspace\nlp\guide\CST 8507_Project_Presentation.pptx"
prs = Presentation(pptx_path)
for i, sl in enumerate(prs.slides):
    for sh in sl.shapes:
        if sh.has_text_frame:
            print(f"Slide {i+1}: {sh.text_frame.text.strip()[:60]}")
            break

print(f"\nFile size: {os.path.getsize(pptx_path) / 1024 / 1024:.1f} MB")
